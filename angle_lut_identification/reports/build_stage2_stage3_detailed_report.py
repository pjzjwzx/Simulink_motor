from __future__ import annotations

import csv
import json
import math
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
PROJECT = ROOT / "angle_lut_identification"
OUT_DIR = PROJECT / "reports"
S2 = PROJECT / "results" / "stage2_20260818_221202_329" / "stage2"
S3 = PROJECT / "results" / "stage3_20260818_234502_642" / "stage3"
DOCX_PATH = OUT_DIR / "Stage2_Stage3_Detailed_Technical_Report_CN.docx"

THEORY_PDF = Path(r"D:\Download\Angle_LUT_Theory_Review_v1.1.pdf")
IMPLEMENTATION_SPEC = Path(r"D:\Download\Codex_Simulink_Angle_LUT_Implementation_Spec.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202B38"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
ORANGE = "D95F0E"
GOLD = "7A5A00"
RISK = "9B1C1C"
WHITE = "FFFFFF"


def load_json(path: Path):
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def load_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def finite(value):
    try:
        x = float(value)
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def pct(value, digits=1):
    x = finite(value)
    return "-" if x is None else f"{100*x:.{digits}f}%"


def num(value, digits=3):
    x = finite(value)
    return "-" if x is None else f"{x:.{digits}f}"


def short_hash(value):
    return f"{value[:12]}...{value[-8:]}" if value else "-"


def read_pair_metrics(root: Path):
    rows = []
    for d in sorted(p for p in root.iterdir() if p.is_dir()):
        p = d / "pair_metrics.json"
        if p.exists():
            rows.append(load_json(p))
    return rows


def set_run_font(run, size=None, bold=None, italic=None, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cant_split)
        for idx, cell in enumerate(row.cells):
            cell.width = Twips(widths_dxa[idx])
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_table_borders(table)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_heading(paragraph, keep_with_next=True):
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = True


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, text, end])


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num_node = OxmlElement("w:num")
    num_node.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num_node.append(abs_id)
    numbering.append(num_node)
    return num_id


def add_bullet(doc, text, num_id):
    p = doc.add_paragraph()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p._p.get_or_add_pPr().append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_para(doc, text="", *, bold_lead=None, size=11, color=INK, align=None, after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size, italic=italic, color=color)
    return p


def add_formula(doc, label, formula, explanation):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F9FC")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{label}  {formula}")
    set_run_font(r, 10.5, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(explanation)
    set_run_font(r2, 9.5, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc, title, body, fill=CALLOUT, title_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, 11, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, 10.5, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths_dxa, aligns=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, font_size, bold=True, color=DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            align = aligns[idx] if aligns else (WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            p.alignment = align
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            color = INK
            bold = False
            if str(value) == "PASS":
                color, bold = DARK_BLUE, True
            elif str(value) == "FAIL":
                color, bold = RISK, True
            set_run_font(r, font_size, bold=bold, color=color)
            set_cell_margins(cells[idx], top=100, bottom=100, start=120, end=120)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_picture(doc, path: Path, caption, source_note, alt_text, width=6.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    set_run_font(r, 9.5, bold=True, color=DARK_BLUE)
    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    src.paragraph_format.space_before = Pt(0)
    src.paragraph_format.space_after = Pt(8)
    r2 = src.add_run(source_note)
    set_run_font(r2, 8.5, italic=True, color=MUTED)


def set_heading_style(style, size, color, before, after):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    set_heading_style(doc.styles["Heading 1"], 16, BLUE, 16, 8)
    set_heading_style(doc.styles["Heading 2"], 13, BLUE, 12, 6)
    set_heading_style(doc.styles["Heading 3"], 12, DARK_BLUE, 8, 4)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    table.autofit = False
    table.columns[0].width = Inches(3.25)
    table.columns[1].width = Inches(3.25)
    table.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for cell, text in zip(table.rows[0].cells, ("ENCODER ANGLE LUT", "STAGE 2 / STAGE 3 TECHNICAL REPORT")):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, 8.5, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("内部技术报告  |  ")
    set_run_font(r, 8.5, color=MUTED)
    add_field(p, "PAGE")


def create_visuals(s2_drift, s3_drift, s2_pairs, s3_pairs):
    def font(size, bold=False):
        candidates = [
            Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
            Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size=size)
        return ImageFont.load_default()

    title_font = font(28, True)
    axis_font = font(19)
    label_font = font(18)
    small_font = font(15)
    legend_font = font(17)
    ink = "#202B38"
    muted = "#667085"
    grid = "#E4E7EC"

    labels = ["-20 rad/s", "-10 rad/s", "-5 rad/s", "+5 rad/s", "+10 rad/s", "+20 rad/s", "Load 0 Nm", "Load 2 Nm"]
    ids = ["speed_-20radps_m", "speed_-10radps_m", "speed_-05radps_m", "speed_+05radps_m", "speed_+10radps_m", "speed_+20radps_m", "load_0Nm", "load_2Nm"]
    s2_map = {r["case_id"]: float(r["lut_drift_e_deg"]) for r in s2_drift}
    s3_map = {r["case_id"]: float(r["lut_drift_e_deg"]) for r in s3_drift}
    v2 = [s2_map[i] for i in ids]
    v3 = [s3_map[i] for i in ids]
    image = Image.new("RGB", (1650, 950), "white")
    draw = ImageDraw.Draw(image)
    left, top, right, bottom = 275, 115, 140, 115
    plot_w = image.width - left - right
    plot_h = image.height - top - bottom
    max_x = 0.86
    draw.text((image.width / 2, 35), "Cross-condition LUT drift: Stage 2 passes; Stage 3 fails at ±5 rad/s", fill=ink, font=title_font, anchor="ma")
    for tick in [0, 0.2, 0.4, 0.6, 0.8]:
        xx = left + plot_w * tick / max_x
        draw.line((xx, top, xx, top + plot_h), fill=grid, width=2)
        draw.text((xx, top + plot_h + 18), f"{tick:.1f}", fill=muted, font=small_font, anchor="ma")
    gate_x = left + plot_w * 0.5 / max_x
    for yy in range(top, top + plot_h, 18):
        draw.line((gate_x, yy, gate_x, min(yy + 9, top + plot_h)), fill="#344054", width=2)
    row_h = plot_h / len(labels)
    bar_h = 25
    for idx, (lab, val2, val3) in enumerate(zip(labels, v2, v3)):
        cy = top + row_h * (idx + 0.5)
        draw.text((left - 18, cy), lab, fill=ink, font=label_font, anchor="rm")
        w2 = plot_w * val2 / max_x
        w3 = plot_w * val3 / max_x
        draw.rectangle((left, cy - 29, left + w2, cy - 29 + bar_h), fill="#2E74B5")
        c3 = "#D95F0E" if val3 >= 0.5 else "#E9A66A"
        draw.rectangle((left, cy + 5, left + w3, cy + 5 + bar_h), fill=c3, outline="#A3410A", width=1)
        draw.text((left + w2 + 10, cy - 16), f"{val2:.3f}", fill="#1F4D78", font=small_font, anchor="lm")
        draw.text((left + w3 + 10, cy + 17), f"{val3:.3f}", fill="#9B1C1C" if val3 >= 0.5 else "#7A3B12", font=small_font, anchor="lm")
    draw.line((left, top + plot_h, left + plot_w, top + plot_h), fill="#344054", width=2)
    draw.line((left, top, left, top + plot_h), fill="#344054", width=2)
    draw.text((left + plot_w / 2, image.height - 35), "Circular LUT drift (deg_e)", fill=ink, font=axis_font, anchor="ma")
    legend_y = 82
    draw.rectangle((left + 20, legend_y - 9, left + 55, legend_y + 9), fill="#2E74B5")
    draw.text((left + 67, legend_y), "Stage 2 / Scheme 4", fill=ink, font=legend_font, anchor="lm")
    draw.rectangle((left + 310, legend_y - 9, left + 345, legend_y + 9), fill="#D95F0E")
    draw.text((left + 357, legend_y), "Stage 3 / Scheme 5", fill=ink, font=legend_font, anchor="lm")
    draw.line((left + 620, legend_y, left + 665, legend_y), fill="#344054", width=2)
    draw.text((left + 677, legend_y), "0.5 deg gate", fill=ink, font=legend_font, anchor="lm")
    drift_path = OUT_DIR / "stage2_stage3_drift_comparison.png"
    image.save(drift_path, quality=95)

    order = ["fixed_00deg_e", "fixed_05deg_e", "fixed_10deg_e", "fixed_20deg_e", "periodic_constant", "periodic_1x", "periodic_2x", "periodic_combined"]
    display = ["Fixed 0°", "Fixed 5°", "Fixed 10°", "Fixed 20°", "Constant", "1×", "2×", "Combined"]

    def pair_map(rows):
        return {r["case_id"]: r for r in rows}

    maps = [pair_map(s2_pairs), pair_map(s3_pairs)]
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    profile_title_font = font(38, True)
    profile_axis_font = font(28)
    profile_label_font = font(27)
    profile_small_font = font(23)
    profile_legend_font = font(24)
    draw.text((900, 30), "Independent frozen pairs converge to the encoder quantization floor", fill=ink, font=profile_title_font, anchor="ma")
    panel_top = 155
    panel_bottom = 115
    panel_lefts = [250, 1125]
    panel_w = 580
    plot_h = image.height - panel_top - panel_bottom
    max_x = 21.4
    for panel_idx, (mapping, panel_title, active_color, panel_left) in enumerate(zip(maps, ["Stage 2 / Scheme 4", "Stage 3 / Scheme 5"], ["#2E74B5", "#D95F0E"], panel_lefts)):
        off = [mapping[i]["baseline"]["control_angle_rmse_e_deg"] for i in order]
        on = [mapping[i]["active"]["control_angle_rmse_e_deg"] for i in order]
        floor_x = panel_left + panel_w * 2.231234 / max_x
        draw.rectangle((panel_left, panel_top, floor_x, panel_top + plot_h), fill="#E8EEF5")
        for tick in [0, 5, 10, 15, 20]:
            xx = panel_left + panel_w * tick / max_x
            draw.line((xx, panel_top, xx, panel_top + plot_h), fill=grid, width=2)
            draw.text((xx, panel_top + plot_h + 20), str(tick), fill=muted, font=profile_small_font, anchor="ma")
        row_h = plot_h / len(order)
        for idx, (lab, off_v, on_v) in enumerate(zip(display, off, on)):
            cy = panel_top + row_h * (idx + 0.5)
            if panel_idx == 0:
                draw.text((panel_left - 18, cy), lab, fill=ink, font=profile_label_font, anchor="rm")
            x_off = panel_left + panel_w * off_v / max_x
            x_on = panel_left + panel_w * on_v / max_x
            draw.line((x_on, cy, x_off, cy), fill="#C7CDD6", width=4)
            draw.ellipse((x_off - 9, cy - 9, x_off + 9, cy + 9), fill="white", outline="#667085", width=3)
            draw.ellipse((x_on - 10, cy - 10, x_on + 10, cy + 10), fill=active_color, outline="white", width=2)
            draw.text((x_on + 14, cy), f"{on_v:.2f}", fill=active_color, font=profile_small_font, anchor="lm")
        draw.line((panel_left, panel_top + plot_h, panel_left + panel_w, panel_top + plot_h), fill="#344054", width=2)
        draw.line((panel_left, panel_top, panel_left, panel_top + plot_h), fill="#344054", width=2)
        draw.text((panel_left + panel_w / 2, 116), panel_title, fill=ink, font=profile_axis_font, anchor="ma")
        draw.text((panel_left + panel_w / 2, image.height - 35), "Control-angle RMSE (deg_e)", fill=ink, font=profile_axis_font, anchor="ma")
    legend_y = 77
    draw.ellipse((510, legend_y - 11, 532, legend_y + 11), fill="white", outline="#667085", width=3)
    draw.text((545, legend_y), "Active off", fill=ink, font=profile_legend_font, anchor="lm")
    draw.ellipse((730, legend_y - 12, 754, legend_y + 12), fill="#D95F0E", outline="white", width=2)
    draw.text((767, legend_y), "Active on", fill=ink, font=profile_legend_font, anchor="lm")
    draw.rectangle((945, legend_y - 11, 985, legend_y + 11), fill="#E8EEF5")
    draw.text((999, legend_y), "Floor +0.10°", fill=ink, font=profile_legend_font, anchor="lm")
    profile_path = OUT_DIR / "stage2_stage3_control_angle_profiles.png"
    image.save(profile_path, quality=95)
    return drift_path, profile_path


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    s2_gate = load_json(S2 / "gate.json")
    s3_gate = load_json(S3 / "gate.json")
    s2_manifest = load_json(S2 / "run_manifest.json")
    s3_manifest = load_json(S3 / "run_manifest.json")
    s2_metrics = load_csv(S2 / "metrics.csv")
    s3_metrics = load_csv(S3 / "metrics.csv")
    s2_train = [r for r in s2_metrics if r["phase"] == "M64"]
    s3_train = [r for r in s3_metrics if r["phase"] == "M64_NOMINAL"]
    s2_drift = load_csv(S2 / "condition_drift" / "metrics.csv")
    s3_drift = [r for r in s3_metrics if r["phase"] == "CONDITION_DRIFT"]
    s2_pairs = read_pair_metrics(S2 / "profile_freeze")
    s3_pairs = read_pair_metrics(S3 / "profile_freeze")
    learning = load_csv(S3 / "reports" / "lut_learning_evolution.csv")
    drift_chart, profile_chart = create_visuals(s2_drift, s3_drift, s2_pairs, s3_pairs)

    doc = Document()
    configure_document(doc)
    bullet_num_id = add_numbering(doc)

    # Cover / masthead
    add_para(doc, "技术验证报告", size=10, color=GOLD, after=4)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("编码器周期角度误差 LUT\nStage 2 与 Stage 3 详细报告")
    set_run_font(r, 24, bold=True, color=DARK_BLUE)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(18)
    r2 = p2.add_run("Scheme 4 流式充分统计量 vs. Scheme 5 二维非线性更新")
    set_run_font(r2, 13.5, color=MUTED)

    meta = [
        ("报告日期", "2026-08-19"),
        ("Stage 2 正式运行", s2_manifest["run_id"]),
        ("Stage 3 正式运行", s3_manifest["run_id"]),
        ("基线", "Stage 1 PASS / 20 kHz / p=21 / N=1024"),
    ]
    add_table(doc, ["项目", "内容"], meta, [1900, 7460], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=10)

    status_rows = [
        ("Stage 2", "PASS", "完整范围：8×M64 + 8×M128 + 8 漂移 + 32 冻结配对"),
        ("Stage 3", "FAIL", "训练收敛，但 ±5 rad/s 的 LUT 漂移超过 0.5°e 门槛并早停"),
    ]
    add_table(doc, ["阶段", "正式状态", "结论"], status_rows, [1500, 1300, 6560], aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT], font_size=10)
    add_callout(
        doc,
        "一句话结论",
        "Stage 2 已证明基于 z 的 Scheme 4 可以稳定辨识并冻结补偿；Stage 3 也证明 Scheme 5 会随融合次数持续改善并在训练工况收敛，但其低速跨工况一致性不足，因此 FAIL 不是“没有学会”，而是“学到的表在 ±5 rad/s 时发生了不可接受的速度相关漂移”。",
        fill="EAF2F8",
    )

    doc.add_page_break()

    doc.add_heading("技术摘要：Stage 2 可交付，Stage 3 需解决低速不变性", level=1)
    add_bullet(doc, f"Stage 2 正式结果为 {s2_gate['status']} / {s2_gate['scope']}。114 项测试全部通过，8 个 M64 与 8 个 M128 训练结果覆盖率均为 100%，32 个冻结配对全部完成。", bullet_num_id)
    add_bullet(doc, f"Stage 2 的主组合波形 M64 active LUT RMSE 为 {s2_gate['aggregate']['primary_m64_active_rmse_e_deg']:.3f}°e，零表到 active 的 LUT RMSE 改善为 96.94%；最大跨工况漂移仅 {s2_gate['aggregate']['maximum_condition_drift_e_deg']:.3f}°e，小于 0.5°e。", bullet_num_id)
    add_bullet(doc, f"Stage 3 的 8 个 M64、8 个 M128、16 个诊断幅值模式和 8 个 profile 冻结配对均完成，117 项测试全通过；主组合波形 active LUT RMSE 为 {s3_gate['aggregate']['primary_m64_active_rmse_e_deg']:.3f}°e。", bullet_num_id)
    add_bullet(doc, "Stage 3 在 condition_drift 阶段失败：-5 rad/s 漂移 0.650°e，+5 rad/s 漂移 0.766°e，均超过 0.5°e。由于 gate 规定失败立即早停，后续敏感度与 24 个完整冻结验证未执行。", bullet_num_id)
    add_bullet(doc, "学习演化证据明确：Scheme 5 active LUT RMSE 从第 1 次融合的 8.426°e 降到第 18 次的 0.438°e；相对零表改善从 15.23% 提高到 95.59%。", bullet_num_id)

    doc.add_heading("阅读导航", level=2)
    nav_rows = [
        ("1", "证据范围与指标定义", "明确哪些结论来自正式 gate、哪些属于诊断推断"),
        ("2", "两种学习方案", "说明 Scheme 4 与 Scheme 5 的输入、更新和控制反馈"),
        ("3", "Gate 修订", "解释 60% 或量化地板分支为什么必要"),
        ("4", "Stage 2 结果", "展示训练、冻结、漂移和工程完整性"),
        ("5", "Stage 3 结果", "展示学习随时间改善、LUT 线条和 Scheme 4/5 对比"),
        ("6", "Stage 3 FAIL 根因", "量化低速漂移、区分已证事实与待验证机制"),
        ("7", "建议与复现", "给出不放宽门槛的下一步实验顺序"),
    ]
    add_table(doc, ["章节", "主题", "作用"], nav_rows, [900, 2600, 5860], aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.5)

    doc.add_heading("1. 证据范围、数据与指标定义", level=1)
    add_para(doc, "本报告只使用正式 Stage 2/3 结果目录中的 gate.json、metrics.csv、result.json、run_manifest.json、训练/冻结指标和自动报告图。真值只用于训练完成后的评价；算法更新器没有读取 truth、plant voltage 或注入误差。", bold_lead="本报告只使用")
    definitions = [
        ("shadow LUT", "学习器当前估计表。它反映最新学习结果，但不直接驱动控制。"),
        ("active LUT", "shadow 经过幅值/单调等约束，并以 gamma=0.2 慢融合后的控制表。控制角使用它。"),
        ("zero LUT", "全零补偿表，即未补偿基线。用于计算 LUT RMSE 改善率。"),
        ("evaluation truth", "注入的真实角误差，只在更新结束后评分，不进入学习或门控。"),
        ("control-angle RMSE", "修正后的控制电角与真实电角之间的 RMS 误差，单位 deg_e。"),
        ("LUT drift", "独立工况学习出的 shadow LUT 与主训练 LUT 的圆周 RMS 差。Stage 2/3 门槛均为 <0.5°e。"),
        ("改善率", "1 - metric_active_on / metric_active_off。正数表示改善；iq tracking change 为变化率，负数表示误差下降。"),
    ]
    add_table(doc, ["术语", "定义"], definitions, [2100, 7260], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.5)
    add_callout(doc, "解释边界", "Stage 3 的 FAIL 是正式的描述性验证结论：两个低速独立 LUT 与主 LUT 的差超过预注册门槛。关于“为什么低速会漂移”的算法机制属于诊断推断；由于早停使 sensitivity_count=0，尚不能把某个单一非理想因素定性为唯一根因。", fill="FFF8E8", title_color=GOLD)

    doc.add_heading("2. 两种学习方案如何工作", level=1)
    doc.add_heading("2.1 Stage 2 / Scheme 4：从 z 流式积累正规方程", level=2)
    add_para(doc, "Scheme 4 的单样本输入为原始机械角 φ、伪测量角 z、质量权重 χ、解缠机械角和时间戳。每个样本只更新两个相邻周期节点对应的固定尺寸充分统计量；不保存样本历史。达到覆盖、权重和累计行程条件后，先用 mldivide 产生参考解，再用固定循环带状求解器求 shadow。")
    add_formula(doc, "E20", "z̃ = unwrap_local(z | hᵀ·shadow)", "以当前 shadow 的局部预测为分支中心，避免 atan2 的 ±π 跳变污染统计量。")
    add_formula(doc, "E21", "A ← A + χhhᵀ,   b ← b + χhz̃", "每次只影响周期线性插值的两个节点；A、b、节点权重、访问次数和行程均为固定尺寸状态。")
    add_formula(doc, "E22", "shadow = arg min ||Hw-z̃||²χ + λs||D²w||² + λ0||w-wprior||²", "shadow 经约束处理后以 gamma=0.2 融入 active，单次 active 节点变化不超过 2°e。")

    doc.add_heading("2.2 Stage 3 / Scheme 5：用二维残差切向分量做非线性局部更新", level=2)
    add_para(doc, "Scheme 5 不再把 atan2 得到的 z 作为学习目标，而是直接使用统一残差 y=[yd,yq]。正式主路固定 Rs/Ls/ψf，采用 nominal_model 幅值；measured_magnitude 与 direction_normalized 仅作诊断，不能在看过真值后替换主路。")
    add_formula(doc, "E30", "δ̂=hᵀ·shadow;  â=Ts·ψf·|ω̂e|/Ls;  ŷ=â[sinδ̂, cosδ̂]ᵀ", "LUT 给出角度 δ̂，名义电机模型给出残差幅值 â。")
    add_formula(doc, "E31", "t=ed·cosδ̂-eq·sinδ̂;  n=ed·sinδ̂+eq·cosδ̂", "t 是沿角度变化方向的切向误差，用于 LUT 更新；n 仅作径向诊断。")
    add_formula(doc, "E32", "w̃←w+μ5·χ·(â·t)/(ε5+â²||h||²)·h", "μ5=0.01，ε5=0.01 A²；正号更新，且单样本只改两个局部 shadow 节点。")

    add_callout(doc, "控制链隔离", "两阶段都保持 raw identification frame 不变：原始机械角/电角进入残差与学习；只有 active LUT 经过周期插值后从控制角中相减。active 关闭时与上一级 harness 数值等价。", fill="EAF2F8")

    doc.add_heading("3. Gate 修订：60% 或量化地板，不是降低所有要求", level=1)
    quant_step = 360 * 21 / 1024
    quant_floor = quant_step / math.sqrt(12)
    floor_limit = quant_floor + 0.1
    add_para(doc, f"编码器 1024 机械计数、21 极对导致每个机械计数对应 {quant_step:.6f}°e。均匀量化的理论 RMS 地板为 {quant_floor:.6f}°e；加 0.10°e 余量后，地板分支阈值为 {floor_limit:.6f}°e。")
    gate_rows = [
        ("理想非零", "改善 ≥60%", "percentage", "达到比例目标"),
        ("理想非零", f"active RMSE ≤{floor_limit:.6f}°e", "quantization_floor", "已触及编码器可解释地板"),
        ("fixed-zero", "不判断百分比", "fixed_zero_exempt", "零误差时分母过小，百分比无意义"),
        ("非理想", "改善 ≥50%", "percentage only", "不允许使用地板例外"),
    ]
    add_table(doc, ["案例类型", "门槛", "记录分支", "目的"], gate_rows, [1700, 2200, 2300, 3160], aligns=[WD_ALIGN_PARAGRAPH.LEFT]*4, font_size=9.5)
    add_para(doc, "该变更是用户授权、版本化且只对新运行生效。旧的 Stage 2 FAIL 结果没有被覆盖或改写；其他求解器、精度、安全、冻结和漂移门槛保持不变。")
    add_callout(doc, "为什么 periodic 2× 可以低于 60% 仍通过", f"Stage 2 的 periodic_2x 改善为 56.80%，Stage 3 为 56.38%，但 active 控制角 RMSE 分别约 2.14°e 和 2.16°e，均低于 {floor_limit:.6f}°e，因此使用 quantization_floor 分支。它们不是被忽略，而是已经达到编码器量化所允许的误差地板。", fill="FFF8E8", title_color=GOLD)

    doc.add_page_break()
    doc.add_heading("4. Stage 2：Scheme 4 完整通过", level=1)
    add_para(doc, "Stage 2 的正式结果是完整范围 PASS。所有训练、求解器、跨工况漂移、冻结控制、安全约束和报告工件均闭合；其结果可作为 Stage 3 的前置输入。", bold_lead="Stage 2 的正式结果是完整范围 PASS。")
    summary2 = [
        ("正式状态", f"{s2_gate['status']} / {s2_gate['scope']}"),
        ("训练规模", f"{s2_gate['aggregate']['m64_profile_count']} 个 M64 + {s2_gate['aggregate']['m128_profile_count']} 个 M128，同流重放"),
        ("冻结验证", f"{s2_gate['aggregate']['frozen_pair_count']} 组 active-off/on（理想 {s2_gate['aggregate']['ideal_pair_count']}，非理想 {s2_gate['aggregate']['nonideal_pair_count']}，饱和 {s2_gate['aggregate']['saturation_pair_count']}）"),
        ("测试", f"{s2_manifest['tests']['count']} PASS / {s2_manifest['tests']['failed']} FAIL"),
        ("求解器最大相对差", f"{s2_gate['aggregate']['maximum_solver_relative_difference']:.3e}（门槛 1e-9）"),
        ("最小 rcond", f"{s2_gate['aggregate']['minimum_rcond']:.3f}（门槛 1e-8）"),
        ("主 M64 active RMSE", f"{s2_gate['aggregate']['primary_m64_active_rmse_e_deg']:.3f}°e"),
        ("最大跨工况漂移", f"{s2_gate['aggregate']['maximum_condition_drift_e_deg']:.3f}°e（门槛 <0.5°e）"),
    ]
    add_table(doc, ["指标", "Stage 2 正式结果"], summary2, [2700, 6660], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.5)

    doc.add_heading("4.1 LUT 对比图：三条线分别代表什么", level=2)
    add_para(doc, "黑色线是注入误差的 evaluation reference，只在训练完成后用于评分；蓝色线是最新 shadow 解；橙色线是经过约束和慢融合、实际可送入控制器的 active 表。黑线与学习器数据流物理隔离，因此蓝/橙线贴近黑线证明的是盲辨识效果，不是真值回灌。")
    add_picture(
        doc,
        S2 / "reports" / "lut_truth_shadow_active.png",
        "图 1  Stage 2 主组合波形的真值、shadow 与 active LUT",
        "来源：Stage 2 正式报告工件；truth 仅用于 evaluation。",
        "Stage 2 LUT 曲线：黑色评价真值、蓝色 shadow、橙色 active，三者在完整机械周期上高度贴合。",
    )
    add_para(doc, "active 与 shadow 不完全重合是预期行为：active 每次只融合 shadow 的 20%，且单节点每次最多变化 2°e。该滞后用于避免控制补偿突跳。最终主组合波形 shadow RMSE 为 0.256°e，active RMSE 为 0.304°e，active 最大节点误差为 0.481°e。")

    doc.add_page_break()
    doc.add_heading("4.2 八条 M64 训练均覆盖完整且稳定求解", level=2)
    rows2_train = []
    for r in s2_train:
        rows2_train.append((
            r["case_id"],
            f"{float(r['shadow_rmse_e_deg']):.3f}",
            f"{float(r['active_rmse_e_deg']):.3f}",
            f"{100*float(r['lut_improvement_fraction']):.2f}%" if float(r["lut_improvement_fraction"]) >= 0 else "fixed-zero",
            r["solve_count"],
            "PASS",
        ))
    add_table(doc, ["训练波形", "shadow RMSE °e", "active RMSE °e", "相对零表改善", "求解次数", "结果"], rows2_train, [2150, 1400, 1400, 1500, 1100, 1810], font_size=8.9)
    add_para(doc, "fixed-zero 的“相对零表改善”为负并不表示算法失稳：零表本来就是真值，任何非零估计都会使相对比例失真。因此 fixed-zero 使用绝对 LUT RMSE 门槛，最终 active RMSE 为 0.253°e，满足 ≤0.5°e。")

    doc.add_heading("4.3 冻结控制：控制角被压到约 2.14–2.26°e 的量化地板", level=2)
    add_para(doc, "每条 LUT 学完后均在独立仿真中做 active-off/on 配对。下图使用同一横轴展示补偿前后控制角 RMSE；浅蓝区为 2.231234°e 的地板阈值。")
    add_picture(
        doc,
        profile_chart,
        "图 2  Stage 2 与 Stage 3 八个 profile 的独立冻结控制角 RMSE",
        "来源：两次正式运行的 profile_freeze/pair_metrics.json；图由本报告重新汇总。",
        "两个面板的哑铃图对比 active 关闭与开启的控制角 RMSE，开启后均收敛到约 2.2 度电角附近。",
    )
    rows2_pair = []
    for r in s2_pairs:
        rows2_pair.append((r["case_id"], f"{r['baseline']['control_angle_rmse_e_deg']:.3f}", f"{r['active']['control_angle_rmse_e_deg']:.3f}", f"{100*r['control_angle_improvement']:.2f}%", r["control_angle_gate_basis"]))
    add_table(doc, ["Stage 2 profile", "off °e", "on °e", "改善", "Gate 分支"], rows2_pair, [2200, 1200, 1200, 1300, 3460], font_size=9.1)

    doc.add_heading("4.4 跨方向、速度与负载 LUT 漂移全部小于 0.5°e", level=2)
    add_para(doc, "Stage 2 的 8 条独立 shadow LUT 与主组合 LUT 的差均低于门槛。最大值出现在 load 2 Nm，为 0.239°e；速度项最大约 0.085°e，说明 Scheme 4 的结果对速度与方向较稳定。")
    add_picture(
        doc,
        S2 / "reports" / "direction_speed_load_drift.png",
        "图 3  Stage 2 的方向、速度和负载一致性",
        "来源：Stage 2 正式报告工件；红线为 0.5°e gate。",
        "Stage 2 八个跨工况漂移柱均明显低于 0.5 度电角门槛。",
    )

    doc.add_page_break()
    doc.add_heading("4.5 Stage 2 的历史失败与修复链", level=2)
    timeline_rows = [
        ("stage2_20260818_194629_441", "FAIL", "旧协议要求理想控制角改善 ≥70%，fixed 5° 仅约 60.05%。结果保留不改。"),
        ("协议 v2", "版本化修订", "用户授权 60% 或量化地板+0.10°e；fixed-zero 豁免，非理想仍为 50%。"),
        ("stage2_20260818_214347_814", "FAIL", "发现 active-on 平衡初始化仍用未补偿控制帧，fixed 20° 配对的平均转矩可比性失真。"),
        ("D014 修复", "实现修复", "冻结 active-on 按实际量化角减 frozen LUT 建立平衡初值；不改控制器、LUT、门槛或评价窗。"),
        (s2_manifest["run_id"], "PASS", "完整训练、漂移和 32 组冻结配对通过。"),
    ]
    add_table(doc, ["节点", "状态", "说明"], timeline_rows, [2550, 1400, 5410], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.1)

    doc.add_heading("5. Stage 3：学习收敛，但跨工况不变性失败", level=1)
    add_para(doc, "Stage 3 的训练阶段不是失败点。八条 nominal M64、八条 M128 同流扫描和两个诊断幅值模式均有限、覆盖完整且满足安全约束；八个 profile 冻结控制也通过。正式 FAIL 发生在随后的 condition_drift gate。", bold_lead="Stage 3 的训练阶段不是失败点。")
    summary3 = [
        ("正式 gate", f"{s3_gate['status']} / {s3_gate['scope']}"),
        ("已完成训练", f"M64 {s3_gate['aggregate']['m64_profile_count']}，M128 {s3_gate['aggregate']['m128_profile_count']}，诊断模式 {s3_gate['aggregate']['diagnostic_profile_count']}"),
        ("已完成冻结", f"profile pair {s3_gate['aggregate']['frozen_pair_count']}；完整验证矩阵因早停未运行"),
        ("测试", f"{s3_manifest['tests']['count']} PASS / {s3_manifest['tests']['failed']} FAIL"),
        ("主 M64 active RMSE", f"{s3_gate['aggregate']['primary_m64_active_rmse_e_deg']:.3f}°e"),
        ("主 M128 active RMSE", f"{s3_gate['aggregate']['primary_m128_active_rmse_e_deg']:.3f}°e"),
        ("最大跨工况漂移", f"{s3_gate['aggregate']['maximum_condition_drift_e_deg']:.3f}°e（门槛 <0.5°e）"),
        ("早停位置", s3_gate["critical_early_stop"]["phase"]),
    ]
    add_table(doc, ["指标", "Stage 3 正式结果"], summary3, [2700, 6660], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.5)

    doc.add_heading("5.1 LUT 确实随学习增加而持续改善", level=2)
    first = learning[0]
    last = learning[-1]
    add_para(doc, f"主组合波形在第 1 次融合时，shadow RMSE 已降到 {float(first['shadow_rmse_e_deg']):.3f}°e，但 active 因慢融合仍为 {float(first['active_rmse_e_deg']):.3f}°e；到第 {last['fusion_event']} 次融合，active RMSE 降到 {float(last['active_rmse_e_deg']):.3f}°e，改善率从 {100*float(first['improvement_fraction']):.2f}% 上升到 {100*float(last['improvement_fraction']):.2f}%。")
    add_picture(
        doc,
        S3 / "reports" / "lut_learning_evolution.png",
        "图 4  Scheme 5 主组合波形的 LUT 学习演化",
        "来源：Stage 3 runner 外部记录的 fusion history；历史仅用于报告，不反馈学习器。",
        "左图显示 shadow 和 active LUT RMSE 随 18 次融合下降；右图显示相对零表改善持续上升并越过 80% gate。",
    )
    add_para(doc, "左图中蓝色 shadow RMSE 在第 6 次融合后大致稳定在 0.39–0.45°e；橙色 active RMSE 因 gamma=0.2 慢融合持续下降，最终追到 0.438°e。右图绿色改善率越过 80% 门槛后继续上升并趋于饱和。这正是“学习越多、补偿越好”的证据。")

    doc.add_heading("5.2 Stage 3 LUT 图的四条线", level=2)
    add_para(doc, "黑色 truth 是评价参考；蓝色 Scheme 5 shadow 是当前学习表；橙红色 Scheme 5 active 是约束和慢融合后的控制表；灰色虚线 zero LUT 是未补偿基线。zero LUT 与 truth 的距离很大，而 shadow/active 与 truth 高度贴合，说明训练工况下学习有效。")
    add_picture(
        doc,
        S3 / "reports" / "lut_truth_shadow_active.png",
        "图 5  Stage 3 主组合波形的 truth / shadow / active / zero LUT",
        "来源：Stage 3 正式报告工件；truth 仅用于训练完成后的评价。",
        "Stage 3 四条 LUT 曲线，学习得到的 shadow 与 active 紧贴评价真值，零表为水平虚线。",
    )

    doc.add_heading("5.3 八条 nominal M64 训练均通过", level=2)
    rows3_train = []
    for r in s3_train:
        improvement = float(r["active_improvement_fraction"])
        rows3_train.append((
            r["case_id"],
            f"{float(r['shadow_rmse_e_deg']):.3f}",
            f"{float(r['active_rmse_e_deg']):.3f}",
            f"{100*improvement:.2f}%" if improvement >= 0 else "fixed-zero",
            r["fusion_count"],
            "PASS",
        ))
    add_table(doc, ["训练波形", "shadow RMSE °e", "active RMSE °e", "相对零表改善", "融合次数", "结果"], rows3_train, [2150, 1400, 1400, 1500, 1100, 1810], font_size=8.9)

    doc.add_heading("5.4 Scheme 5 与 Scheme 4 的最终 LUT 非常接近", level=2)
    add_para(doc, "左面板比较 shadow，右面板比较 active：蓝色虚线为 Scheme 4，橙色实线为 Scheme 5。两种算法虽然目标不同（Scheme 4 学 z，Scheme 5 学二维 y），但主训练工况下得到的表高度一致。")
    add_picture(
        doc,
        S3 / "reports" / "scheme4_scheme5_comparison.png",
        "图 6  同一训练流下 Scheme 4 与 Scheme 5 的 shadow / active 对比",
        "来源：Stage 3 正式报告工件；Scheme 4 表来自通过的 Stage 2。",
        "两个面板显示 Scheme 4 蓝色虚线与 Scheme 5 橙色实线几乎重合。",
    )
    add_para(doc, "主组合 M64 的 Scheme 5 与 Scheme 4 shadow 圆周 RMS 差为 0.168°e，active 差为 0.173°e，均低于 0.5°e。由此可排除“Stage 3 完全学错方向或符号”的解释。")

    doc.add_heading("5.5 三种幅值模式：正式主路不是数值最小者", level=2)
    add_para(doc, "measured_magnitude 与 direction_normalized 在主组合波形上的 active RMSE 约 0.35°e，低于 nominal_model 的 0.438°e；但协议预注册 nominal_model 为唯一 gate-driving 路径，因此没有依据真值改选模式。这保护了结果免受事后择优。")
    add_picture(
        doc,
        S3 / "reports" / "amplitude_mode_comparison.png",
        "图 7  Stage 3 三种幅值模式的 active LUT RMSE",
        "来源：Stage 3 正式报告工件；后两种仅诊断，不决定最终 gate。",
        "柱状图显示 direction_normalized 与 measured_magnitude 约 0.35 度，nominal_model 约 0.44 度。",
    )

    doc.add_heading("6. Stage 3 为什么 FAIL", level=1)
    add_para(doc, "直接原因只有一个：独立低速 LUT 与主训练 LUT 的圆周 RMS 差超过 0.5°e。-5 rad/s 为 0.650°e，超门槛约 30%；+5 rad/s 为 0.766°e，超门槛约 53%。其余六个方向/速度/负载案例通过。", bold_lead="直接原因只有一个：")
    add_picture(
        doc,
        drift_chart,
        "图 8  Stage 2 与 Stage 3 的跨工况 LUT 漂移对比",
        "来源：两次正式运行的 condition-drift 指标；图由本报告重新汇总。",
        "水平分组条形图显示 Stage 2 八项均小于 0.5 度，而 Stage 3 的正负 5 rad/s 两项越过门槛。",
    )
    s2_map = {r["case_id"]: float(r["lut_drift_e_deg"]) for r in s2_drift}
    s3_map = {r["case_id"]: float(r["lut_drift_e_deg"]) for r in s3_drift}
    drift_order = ["speed_-20radps_m", "speed_-10radps_m", "speed_-05radps_m", "speed_+05radps_m", "speed_+10radps_m", "speed_+20radps_m", "load_0Nm", "load_2Nm"]
    drift_rows = []
    for cid in drift_order:
        verdict = "FAIL" if s3_map[cid] >= 0.5 else "PASS"
        drift_rows.append((cid, f"{s2_map[cid]:.3f}", f"{s3_map[cid]:.3f}", "<0.500", verdict))
    add_table(doc, ["独立工况", "Stage 2 °e", "Stage 3 °e", "门槛 °e", "Stage 3"], drift_rows, [2450, 1450, 1450, 1400, 2610], font_size=9.1)

    doc.add_page_break()
    doc.add_heading("6.1 这不是训练精度、符号或安全约束失败", level=2)
    evidence_rows = [
        ("训练精度", "通过", "主 M64 shadow/active RMSE=0.419/0.438°e；非零 profile LUT 改善均 >90%。"),
        ("学习随时间收敛", "通过", "active RMSE 8.426→0.438°e，改善 15.23%→95.59%。"),
        ("Scheme 4 一致性", "通过", "主表 shadow/active 差约 0.168/0.173°e。"),
        ("冻结控制", "通过", "8 个 profile pair 全部通过 percentage、floor 或 fixed-zero 分支。"),
        ("覆盖与安全", "通过", "M64/M128 覆盖 100%，融合次数、安全约束、未覆盖节点保护均通过。"),
        ("跨工况不变性", "失败", "±5 rad/s 的独立 LUT 漂移为 0.650/0.766°e。"),
    ]
    add_table(doc, ["候选解释", "判定", "正式证据"], evidence_rows, [2000, 1300, 6060], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.3)

    doc.add_heading("6.2 最可能的机制，但尚未被 sensitivity 阶段证实", level=2)
    add_bullet(doc, "E30 的名义幅值 â 与 |ωe| 成正比。低速时二维残差的有效幅值更小，更新更容易受 ADC 量化、电流噪声、时序和模型误差的相对占比影响。", bullet_num_id)
    add_bullet(doc, "±5 rad/s 只有 5–6 次融合，而 ±20 rad/s 有 11 次；融合数量与信息结构不同，可能使独立表对有限样本和局部残差偏差更敏感。", bullet_num_id)
    add_bullet(doc, "E30 使用 |ωe|，理论上正反转更新应一致；但 +5 与 -5 的漂移都失败且幅值不同，提示残差重构、方向门控或非理想项仍有速度/方向相关成分。", bullet_num_id)
    add_bullet(doc, "这些是基于正式数据的诊断推断，不是唯一归因。Stage 3 在漂移 gate 后按合同早停，sensitivity_count=0，因此尚无 Rs/Ls/ψf、死区、延迟等分项证据来锁定根因。", bullet_num_id)
    add_callout(doc, "为什么不能把 Stage 3 改判为 PASS", "Stage 3 的完整范围要求漂移、敏感度、24 组完整冻结矩阵和饱和诊断全部完成并通过。当前 condition_drift 已有两项硬失败，随后阶段依法未执行；即使训练曲线非常好，也不能用训练工况的收敛替代跨工况鲁棒性。", fill="FDECEC", title_color=RISK)

    doc.add_heading("7. 下一步建议：先验证低速机制，再决定是否重开协议", level=1)
    recommendations = [
        ("1", "保持当前 FAIL 工件不变", "不得改写 gate.json、重标案例或放宽 0.5°e 门槛。"),
        ("2", "建立低速定向诊断", "只读重放 ±5 与 ±10 rad/s 的 y、t、n、â、χ、节点访问与融合事件，比较每节点信息量和创新偏置。"),
        ("3", "做等行程/等信息量对照", "分别控制机械周期、有效权重与融合次数，区分“低速本身”与“触发次数不同”。"),
        ("4", "预注册幅值模型 A/B", "若新协议允许，可在不看真值选优的前提下预注册 nominal 与 direction-normalized 的低速一致性对照。"),
        ("5", "验证方向与时序敏感度", "在相同随机种子下扫描正反转、5/50 µs 对齐和电压重构，检查 t/n 与 LUT 漂移是否同向。"),
        ("6", "仅在根因闭合后正式重跑", "新 Stage 3 必须从干净会话完整执行，不能拼接旧训练与新冻结结果形成 PASS。"),
    ]
    add_table(doc, ["顺序", "动作", "验收目的"], recommendations, [900, 2900, 5560], aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.4)

    doc.add_heading("8. 完整性、可复现性与限制", level=1)
    integrity_rows = [
        ("随机种子", str(s2_manifest["random_seed"]), str(s3_manifest["random_seed"])),
        ("测试", f"{s2_manifest['tests']['count']} PASS / 0 FAIL", f"{s3_manifest['tests']['count']} PASS / 0 FAIL"),
        ("源模型 SHA256", short_hash(s2_manifest["source_model_sha256_after"]), short_hash(s3_manifest["source_model_sha256_after"])),
        ("Stage 2 harness", short_hash(s2_manifest["stage2_harness_sha256_after"]), short_hash(s3_manifest["stage2_harness_sha256_after"])),
        ("Stage 3 harness", "-", short_hash(s3_manifest["stage3_harness_sha256"])),
        ("报告工件", "generated", "generated"),
        ("Stage 4", "未执行", "未执行"),
    ]
    add_table(doc, ["审计项", "Stage 2", "Stage 3"], integrity_rows, [2500, 3430, 3430], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9.2)
    add_para(doc, "两次正式运行前后源模型 SHA256 均为 0ceef284...586c2cf，说明原始 SLX 未被修改。Stage 3 的 gate scope 为 BOUNDED_OR_PARTIAL_STAGE3，是早停的结果；run manifest 的 requested scope 仍为 FULL_STAGE3。")
    add_bullet(doc, "本报告没有把 evaluation truth 当作可部署信号；所有 truth 曲线均明确标为 evaluation only。", bullet_num_id)
    add_bullet(doc, "Stage 2 的完整 PASS 证明 Scheme 4 适用于当前案例矩阵，但不等于对真实硬件误差来源的外推保证。", bullet_num_id)
    add_bullet(doc, "Stage 3 低速机制尚未由 sensitivity 案例拆分；报告只把其列为待验证假设。", bullet_num_id)
    add_bullet(doc, "control-angle RMSE 的约 2.13°e 地板来自理想均匀量化假设；真实编码器还可能存在非均匀量化、抖动和时间戳误差。", bullet_num_id)

    doc.add_heading("附录 A：正式工件与复现入口", level=1)
    artifacts = [
        ("Stage 2 gate", str(S2 / "gate.json")),
        ("Stage 2 metrics", str(S2 / "metrics.csv")),
        ("Stage 2 result", str(S2 / "result.json")),
        ("Stage 3 gate", str(S3 / "gate.json")),
        ("Stage 3 metrics", str(S3 / "metrics.csv")),
        ("Stage 3 result", str(S3 / "result.json")),
        ("Stage 3 学习历史", str(S3 / "reports" / "lut_learning_evolution.csv")),
        ("协议决策", str(PROJECT / "docs" / "decisions.md")),
    ]
    add_table(doc, ["工件", "绝对路径"], artifacts, [2200, 7160], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=7.8)
    add_para(doc, "干净会话复现顺序：setup_project; r2=run_stage2; assert(r2.final_status==\"PASS\"); r3=run_stage3;。当前确定性运行预期 Stage 2 PASS、Stage 3 在 condition_drift 返回 FAIL。", size=9.5)

    doc.add_heading("附录 B：来源与版本", level=1)
    source_rows = [
        ("数学/公式", str(THEORY_PDF)),
        ("工程/门控合同", str(IMPLEMENTATION_SPEC)),
        ("MATLAB", "R2024b Update 7 / 24.2.0.3070828"),
        ("Simulink", "24.2 (R2024b)"),
        ("报告证据截止", "Stage 2 2026-08-18 14:59:52；Stage 3 2026-08-18 16:09:10"),
    ]
    add_table(doc, ["来源", "标识"], source_rows, [2200, 7160], aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], font_size=9)

    add_callout(doc, "最终判定", "Stage 2：可接受并可作为后续阶段的稳定基线。Stage 3：学习机制与训练工况性能已得到正面验证，但低速跨工况一致性未达到预注册要求；应保留 FAIL，完成定向低速诊断后再决定是否建立新的 Stage 3 协议版本。", fill="EAF2F8")

    # Stable metadata
    props = doc.core_properties
    props.title = "编码器周期角度误差 LUT - Stage 2 与 Stage 3 详细技术报告"
    props.subject = "Scheme 4 / Scheme 5 implementation, convergence, frozen validation, and Stage 3 failure analysis"
    props.author = "Codex"
    props.keywords = "Simulink, encoder angle error, LUT, Scheme 4, Scheme 5, Stage 2, Stage 3"
    props.created = datetime(2026, 8, 19)
    props.modified = datetime(2026, 8, 19)

    doc.save(DOCX_PATH)

    chart_map = {
        "delivery_surface": "DOCX/PDF technical report",
        "palette_policy": "hard two-root cap plus neutrals",
        "charts": [
            {"section": "Stage 2 LUT", "question": "Does the learned LUT match evaluation truth?", "family": "trend", "type": "multi-series line", "source": str(S2 / "reports" / "lut_truth_shadow_active.png")},
            {"section": "Frozen profiles", "question": "Does active compensation reduce control-angle RMSE to the quantization floor?", "family": "comparison", "type": "two-panel dumbbell", "source": str(profile_chart)},
            {"section": "Stage 2 drift", "question": "Is Scheme 4 condition-invariant?", "family": "comparison", "type": "bar with threshold", "source": str(S2 / "reports" / "direction_speed_load_drift.png")},
            {"section": "Stage 3 learning", "question": "Does LUT accuracy improve with fusion events?", "family": "trend", "type": "two-panel line", "source": str(S3 / "reports" / "lut_learning_evolution.png")},
            {"section": "Stage 3 LUT", "question": "Does Scheme 5 match truth and improve over zero LUT?", "family": "trend", "type": "multi-series line", "source": str(S3 / "reports" / "lut_truth_shadow_active.png")},
            {"section": "Scheme comparison", "question": "Do Scheme 4 and Scheme 5 converge to similar tables?", "family": "comparison", "type": "two-panel line", "source": str(S3 / "reports" / "scheme4_scheme5_comparison.png")},
            {"section": "Amplitude modes", "question": "How do preregistered amplitude modes compare?", "family": "comparison", "type": "bar", "source": str(S3 / "reports" / "amplitude_mode_comparison.png")},
            {"section": "Failure analysis", "question": "Where does Stage 3 exceed the drift gate relative to Stage 2?", "family": "comparison", "type": "grouped horizontal bar", "source": str(drift_chart)},
        ],
    }
    with (OUT_DIR / "stage2_stage3_chart_map.json").open("w", encoding="utf-8") as f:
        json.dump(chart_map, f, ensure_ascii=False, indent=2)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
